#!/usr/bin/env python3
"""QA automático de la presentación (sale con código 1 si algo falla).

Uso: python tesis/presentacion/scripts/qa.py [archivo.pptx]
Requiere: python-pptx, opencv-python-headless (decodificar el QR).
Medidas en pt del lienzo (1920 x 1080).
"""
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation

PRES = Path(__file__).resolve().parent.parent
FILE = Path(sys.argv[1]) if len(sys.argv) > 1 else PRES / 'UniWheels_Semana_Ingenieria_2026.pptx'
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
PT = 12700
VISIBLE = 15            # diapositivas 1-15 visibles; 16+ son respaldo (ocultas)
MIN_SZ, MIN_SZ_BACKUP = 2800, 2600
X0, X1, Y0, Y1 = 86, 1834, 0, 1080
FORBIDDEN = re.compile(r'lorem|ipsum|título\b|subtítulo|pie de foto|para ti', re.I)
QR_URL = 'https://uniwheels.org'

errors = []


def fail(msg):
    errors.append(msg)
    print('FALLA:', msg)


prs = Presentation(str(FILE))

# 1. tamaños de fuente
with zipfile.ZipFile(FILE) as z:
    for i in range(1, len(prs.slides) + 1):
        pass
    for i, slide in enumerate(prs.slides, 1):
        sizes = [int(v) for v in slide._element.xpath('.//@sz')]
        lim = MIN_SZ if i <= VISIBLE else MIN_SZ_BACKUP
        if sizes and min(sizes) < lim:
            fail(f'diapositiva {i}: sz mínimo {min(sizes)} < {lim}')

# 1b. sin negrita sintética sobre ExtraBold
for i, slide in enumerate(prs.slides, 1):
    for r in slide._element.xpath('.//a:rPr[@b="1"]'):
        lat = r.find('a:latin', NS)
        if lat is not None and 'ExtraBold' in lat.get('typeface', ''):
            fail(f'diapositiva {i}: negrita sintética sobre ExtraBold')

# 2 y 3. límites y solapes de shapes con texto
def has_text(sh):
    if sh.has_text_frame and sh.text_frame.text.strip():
        return True
    return sh.shape_type == 19  # tabla


def box(sh):
    return (sh.left / PT, sh.top / PT, (sh.left + sh.width) / PT, (sh.top + sh.height) / PT)


def inside(a, b, tol=1):
    return a[0] >= b[0] - tol and a[1] >= b[1] - tol and a[2] <= b[2] + tol and a[3] <= b[3] + tol


for i, slide in enumerate(prs.slides, 1):
    if i > VISIBLE:
        continue
    items = [(sh.name, box(sh)) for sh in slide.shapes if has_text(sh)]
    for name, b in items:
        if not inside(b, (X0, Y0, X1, Y1)):
            fail(f'diapositiva {i}: «{name}» fuera de {X0}-{X1} x {Y0}-{Y1}: {tuple(round(v) for v in b)}')
    for a in range(len(items)):
        for c in range(a + 1, len(items)):
            na, ba = items[a]
            nc, bc = items[c]
            ix = min(ba[2], bc[2]) - max(ba[0], bc[0])
            iy = min(ba[3], bc[3]) - max(ba[1], bc[1])
            if ix > 0.5 and iy > 0.5 and not (inside(ba, bc) or inside(bc, ba)):
                fail(f'diapositiva {i}: «{na}» se solapa con «{nc}» ({ix:.0f} x {iy:.0f} pt)')

# 4. texto prohibido (diapositivas y notas)
for i, slide in enumerate(prs.slides, 1):
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    if slide.has_notes_slide:
        texts.append(slide.notes_slide.notes_text_frame.text)
    for t in texts:
        m = FORBIDDEN.search(t)
        if m:
            fail(f'diapositiva {i}: texto prohibido «{m.group(0)}»')

# 5. tiempos de las notas
total = 0
for i, slide in enumerate(prs.slides, 1):
    if i > VISIBLE or not slide.has_notes_slide:
        continue
    m = re.match(r'\s*\[(\d+):(\d+)\]', slide.notes_slide.notes_text_frame.text)
    if m:
        total += int(m[1]) * 60 + int(m[2])
print(f'suma de tiempos: {total // 60}:{total % 60:02d}')
if not 9 * 60 + 30 <= total <= 9 * 60 + 50:
    fail(f'suma de tiempos {total // 60}:{total % 60:02d} fuera de 9:30-9:50')

# 6. QR
try:
    import cv2
    qr_slide = prs.slides[13]
    data = None
    for sh in qr_slide.shapes:
        if sh.shape_type == 13 and 'QR' in sh.name:
            img = cv2.imdecode(__import__('numpy').frombuffer(sh.image.blob, 'uint8'), cv2.IMREAD_COLOR)
            data, _, _ = cv2.QRCodeDetector().detectAndDecode(img)
            w_pt = sh.width / PT
            if w_pt < 520:
                fail(f'QR de {w_pt:.0f} pt < 520 pt')
    if data != QR_URL:
        fail(f'el QR decodifica «{data}», se esperaba {QR_URL}')
    else:
        print('QR ->', data)
except Exception as e:  # sin cv2 no se puede verificar: es un fallo
    fail(f'no se pudo decodificar el QR: {e}')

print('\nOK: sin fallas' if not errors else f'\n{len(errors)} falla(s)')
sys.exit(1 if errors else 0)
